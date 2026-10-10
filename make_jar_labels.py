#!/usr/bin/env python3
"""
make_jar_labels.py -- printer-ready jar labels for the CCBER herpetology collection.

Reads the jar-inventory workbook (one row per jar) and writes one PDF with one
label per page, laid out like Label_format_example.pdf:

    Xantusia vigilis   (Xantusiidae)                Jar#: H1175
    U.S.A., California, San Bernardino
    ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    Catalog #        Locality                     Collector
    ────────────────────────────────────────────────────────────
    41719            Antelope Valley              K. Seltmann
    ...
    ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
          UC Santa Barbara Vertebrate Zoology – HERPETOLOGY
      Cheadle Center for Biodiversity and Ecological Restoration

Printing geometry
-----------------
The label prints left-to-right along the roll.  Every page is
  height = the paper width (default 10 cm, --paper-width), and
  length = whatever the content needs (or a fixed --length for a given jar).
If a jar has more specimens than fit down one column, the table continues in
a second (third, ...) column group to the right, so the label gets longer,
never taller.  Use --rotate if your printer driver is set to portrait.

Sizing labels to a jar
----------------------
Give a jar size and the label is cut to that jar instead of to its content:
  height = 60% of the jar's height (capped at the roll width)
  length = --wrap x circumference (default 0.75, i.e. three-quarters around)
Type shrinks (down to --min-scale, default 60%) until the label fits; jars
that still don't fit are listed in the review file.  When the label is
narrower than the roll, a dashed trim line marks its bottom edge.
  --jar-size 250ml          every label for a 250 mL jar
  (no --jar-size)           each row's "Jar Size" cell; blank rows fit to content
  --list-jar-sizes          show the table (edit JAR_SIZES / JAR_ALIASES below)
Accepted sizes: 125ml 250ml 500ml 950ml 2000ml 4000ml, aliases such as 8oz,
quart, gallon, small/medium/large, a Thermo catalog number (121-0250), or any
jar as DIAMETERxHEIGHT in mm (70x150).

Columns used (header names matched case-insensitively)
-------------------------------------------------------
  Jar #, Catalog #s (falls back to Side/Top Catalog), State etc.,
  County/Locality, Scientific Name, Preservative            -- current sheet
  Family, Collector                                          -- optional; add
      them to the sheet and they are used.  Without a Family column the
      family is looked up from the genus via GBIF and cached in
      family_cache.json next to this script.  Without a Collector column the
      Collector column is left off the label.

Font
----
The example uses Avenir.  ReportLab can only embed TrueType (.ttf/.ttc)
fonts, so the script looks for Avenir in the usual macOS locations; otherwise
point --font-regular / --font-italic at TrueType files, or it falls back to
Helvetica (built in, prints well on thermal heads).

Outputs
-------
  <output>.pdf          the labels
  <output>_review.csv   jars whose data needed a judgement call

Requirements:  pip install openpyxl reportlab

Examples
--------
  python make_jar_labels.py Herp_Label_Data.xlsx
  python make_jar_labels.py Herp_Label_Data.xlsx --jar H1874 H1875 -o test.pdf
  python make_jar_labels.py Herp_Label_Data.xlsx --sheet "12,14" --length 18
"""
from __future__ import annotations

import argparse
import csv
import difflib
import json
import math
import re
import sys
import time
import urllib.parse
import urllib.request
from collections import defaultdict
from dataclasses import dataclass, field, replace
from pathlib import Path

import openpyxl
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfbase.ttfonts import TTFont, TTFontFile
from reportlab.pdfgen import canvas

DEFAULT_FOOTER = [
    "UC Santa Barbara Vertebrate Zoology – HERPETOLOGY",
    "Cheadle Center for Biodiversity and Ecological Restoration",
]

# Column headers as they appear in the workbook (lower-case, trimmed).
COL = {
    "jar": ["jar #", "jar#"],
    "catalog": ["catalog #s", "catalog #", "catalog"],
    "top": ["top catalog"],
    "side": ["side catalog"],
    "state": ["state etc.", "state"],
    "locality": ["county/locality", "locality"],
    "sciname": ["scientific name"],
    "family": ["family"],
    "collector": ["collector", "collector(s)", "collectors"],
    "preservative": ["preservative"],
    "shelf": ["shelf #letter"],
    "size": ["jar size"],
}

# ==========================================================================
# Fonts
# ==========================================================================
FONT_REG = "Helvetica"
FONT_ITAL = "Helvetica-Oblique"

AVENIR_CANDIDATES = [
    "/System/Library/Fonts/Avenir.ttc",
    "/System/Library/Fonts/Supplemental/Avenir.ttc",
    "/Library/Fonts/Avenir.ttc",
    str(Path.home() / "Library/Fonts/Avenir.ttc"),
]
AVENIR_FACES = {"regular": [b"Avenir-Book", b"AvenirLTPro-Book", b"Avenir-Roman"],
                "italic": [b"Avenir-BookOblique", b"AvenirLTPro-BookOblique",
                           b"Avenir-Oblique"]}


def _register(alias: str, path: str, index: int = 0) -> str:
    pdfmetrics.registerFont(TTFont(alias, path, subfontIndex=index))
    return alias


def setup_fonts(regular: str | None, italic: str | None) -> str:
    """Register fonts; returns a description of what is being used."""
    global FONT_REG, FONT_ITAL
    if regular:
        FONT_REG = _register("LabelRegular", regular)
        FONT_ITAL = _register("LabelItalic", italic) if italic else FONT_REG
        return f"{Path(regular).name} / {Path(italic).name if italic else '(no italic)'}"

    for path in AVENIR_CANDIDATES:
        if not Path(path).exists():
            continue
        faces = {}
        for i in range(32):
            try:
                faces[TTFontFile(path, subfontIndex=i).name] = i
            except Exception:
                break
        reg = next((faces[n] for n in AVENIR_FACES["regular"] if n in faces), None)
        ita = next((faces[n] for n in AVENIR_FACES["italic"] if n in faces), None)
        if reg is not None and ita is not None:
            FONT_REG = _register("LabelRegular", path, reg)
            FONT_ITAL = _register("LabelItalic", path, ita)
            return f"Avenir ({path})"
    return "Helvetica (Avenir not found; see --font-regular)"


# ==========================================================================
# Text clean-up
# ==========================================================================
US_STATES = [
    "Alabama", "Alaska", "Arizona", "Arkansas", "California", "Colorado",
    "Connecticut", "Delaware", "Florida", "Georgia", "Hawaii", "Idaho",
    "Illinois", "Indiana", "Iowa", "Kansas", "Kentucky", "Louisiana", "Maine",
    "Maryland", "Massachusetts", "Michigan", "Minnesota", "Mississippi",
    "Missouri", "Montana", "Nebraska", "Nevada", "New Hampshire", "New Jersey",
    "New Mexico", "New York", "North Carolina", "North Dakota", "Ohio",
    "Oklahoma", "Oregon", "Pennsylvania", "Rhode Island", "South Carolina",
    "South Dakota", "Tennessee", "Texas", "Utah", "Vermont", "Virginia",
    "Washington", "West Virginia", "Wisconsin", "Wyoming",
]
US_ABBR = {"CA": "California", "AZ": "Arizona", "NM": "New Mexico",
           "NV": "Nevada", "TX": "Texas", "OR": "Oregon", "FL": "Florida"}
AUS_STATES = ["Western Australia", "Northern Territory", "Queensland",
              "New South Wales", "Victoria", "South Australia", "Tasmania"]
COUNTRIES = ["Mexico", "Costa Rica", "France", "Australia", "Canada",
             "Panama", "Ecuador", "Brazil", "Madagascar", "South Africa"]
NULL_GEO = {"", "?", "n/a", "na", "no data", "none", "unlisted", "unknown",
            "out of state", "(out of state)"}
COUNTRY_US = "U.S.A."


def clean(v) -> str:
    """Cell value -> trimmed string ('' for empty). 41261.0 -> '41261'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s.strip('"').strip()


def _fuzzy(name: str, choices: list[str], cutoff=0.85) -> str | None:
    hit = difflib.get_close_matches(name.title(), choices, n=1, cutoff=cutoff)
    return hit[0] if hit else None


def geography_header(state_raw: str, county_raw: str) -> str:
    """'CALIFORNIA' + 'San Bernardino Co.' -> 'U.S.A., California, San Bernardino'."""
    parts: list[str] = []
    st = clean(state_raw)
    if st.lower() in NULL_GEO or st.lower().startswith("mixed"):
        st = ""
    if st:
        if st.upper() in US_ABBR:
            parts = [COUNTRY_US, US_ABBR[st.upper()]]
        elif (m := _fuzzy(st, US_STATES)):
            parts = [COUNTRY_US, m]
        elif "australia" in st.lower() or _fuzzy(st, AUS_STATES):
            region = re.sub(r",?\s*australia\s*$", "", st, flags=re.I).strip(", ")
            region = _fuzzy(region, AUS_STATES) or (region.title() if region else "")
            parts = ["Australia"] + ([region] if region and region != "Australia" else [])
        elif (m := _fuzzy(st, COUNTRIES)):
            parts = [m]
        else:
            parts = [st.title() if st.isupper() else st]

    cty = clean(county_raw)
    m = re.fullmatch(r"([A-Za-z .'-]+?)\s+(Co\.?|County)", cty, flags=re.I)
    if m and parts and parts[0] == COUNTRY_US:
        name = m.group(1).strip()
        parts.append(name.title() if name.isupper() else name)
    return ", ".join(parts)


def locality_text(county_raw: str) -> str:
    s = clean(county_raw)
    return "" if s.lower() in NULL_GEO else s


def normalize_sciname(s: str) -> str:
    """'Bufo Microscaphus ' -> 'Bufo microscaphus' (genus capitalised only)."""
    s = clean(s)
    if not s:
        return ""
    words = s.split(" ")
    out = [words[0][:1].upper() + words[0][1:].lower()]
    out += [w if w.startswith("(") else w.lower() for w in words[1:]]
    return " ".join(out)


ROMAN_TOKENS = {"sp.", "spp.", "sp", "spp", "cf.", "cf", "aff.", "aff", "x", "?",
                "var.", "ssp.", "subsp."}


# ==========================================================================
# Catalog-number parsing
# ==========================================================================
NUM = re.compile(r"^\d+$")
RANGE = re.compile(r"^(\d+)\s*[-–—]\s*(\d+)$")


def parse_catalog(raw: str, max_expand: int, expand=True) -> tuple[list[str], list[str]]:
    """Return (label rows, problems).  Numbers are sorted ascending; anything
    unparseable is kept verbatim at the end and reported."""
    nums: set[int] = set()
    verbatim: list[str] = []
    problems: list[str] = []
    ranges_kept: list[tuple[int, str]] = []

    tokens = []
    for t in re.split(r"[,;]", raw):
        t = t.strip()
        if not t:
            continue
        pieces = t.split()  # "167  16575-16578" -> two tokens
        if len(pieces) > 1 and all(NUM.match(p) or RANGE.match(p) for p in pieces):
            tokens.extend(pieces)
        else:
            tokens.append(t)

    for t in tokens:
        if NUM.match(t):
            nums.add(int(t))
            continue
        m = RANGE.match(t)
        if m:
            a_s, b_s = m.groups()
            a = int(a_s)
            # shorthand: 41961-67 -> 41961-41967
            b = int(a_s[: len(a_s) - len(b_s)] + b_s) if len(b_s) < len(a_s) else int(b_s)
            if b < a:
                verbatim.append(t)
                problems.append(f"reversed range '{t}' printed as-is")
            elif b - a + 1 > max_expand:
                verbatim.append(t)
                problems.append(f"range '{t}' spans {b - a + 1} numbers (> {max_expand}); "
                                "printed as a range")
            elif expand:
                nums.update(range(a, b + 1))
            else:
                ranges_kept.append((a, f"{a}–{b}"))
            continue
        verbatim.append(t)
        problems.append(f"unparsed catalog entry '{t}' printed as-is")

    rows = [str(n) for n in sorted(nums)]
    if ranges_kept:
        rows = [s for _, s in sorted([(n, str(n)) for n in nums] + ranges_kept)]
    return rows + verbatim, problems


RANGE_SEP = " - "


def collapse_runs(rows: list[str], min_run: int = 2, sep: str = RANGE_SEP) -> list[str]:
    """Fold consecutive catalog numbers into ranges.
    ['40748', ..., '40782', '40800'] -> ['40748 - 40782', '40800'].
    Runs shorter than min_run stay as single rows; non-numeric rows pass through."""
    out: list[str] = []
    run: list[int] = []

    def flush():
        if len(run) >= max(min_run, 2):
            out.append(f"{run[0]}{sep}{run[-1]}")
        else:
            out.extend(str(n) for n in run)
        run.clear()

    for r in rows:
        if NUM.match(r):
            n = int(r)
            if run and n != run[-1] + 1:
                flush()
            run.append(n)
        else:
            flush()
            out.append(r)
    flush()
    return out


def looks_like_catalog(s: str) -> bool:
    return bool(s) and bool(re.search(r"\d{3,}", s)) and not re.search(
        r"no (catalog|category|numbers)", s, re.I)


# ==========================================================================
# Family lookup (GBIF backbone, cached)
# ==========================================================================
# Genera GBIF can't resolve on name alone (homonyms in other kingdoms/phyla)
# or misspellings it doesn't catch.  Edit freely; these win over GBIF.
FAMILY_OVERRIDES = {
    "Heterodon": "Colubridae",
    "Micrurus": "Elapidae",
    "Echiopsis": "Elapidae",
    "Eridiphas": "Colubridae",
    "Crolaus": "Viperidae",       # sheet typo for Crotalus
    "Carallus": "Boidae",         # sheet typo for Corallus
    "Ambyiroma": "Ambystomatidae",  # sheet typo for Ambystoma
}
NOT_A_GENUS = {"None", "Unlisted", "Unknown", "Data", "No"}


class FamilyLookup:
    URL = "https://api.gbif.org/v1/species/match?"

    def __init__(self, cache_path: Path, online: bool):
        self.path = cache_path
        self.online = online
        self.cache: dict = {}
        self.dirty = False
        self.failures = 0
        if cache_path.exists():
            try:
                self.cache = json.loads(cache_path.read_text())
            except Exception:
                self.cache = {}

    @staticmethod
    def genus_of(sciname: str) -> str:
        g = re.sub(r"[^A-Za-z]", "", sciname.split(" ")[0]) if sciname else ""
        return g.capitalize() if len(g) > 2 else ""

    def get(self, sciname: str) -> tuple[str, str]:
        """Return (family, note)."""
        genus = self.genus_of(sciname)
        if not genus:
            return "", "no genus to look up family"
        if not genus or genus in NOT_A_GENUS:
            return "", "no genus to look up family"
        if genus.endswith("idae"):          # sheet gave a family, not a genus
            return genus, ""
        if genus in FAMILY_OVERRIDES:
            return FAMILY_OVERRIDES[genus], ""
        if genus not in self.cache and self.online:
            hit = self._query(genus)
            if hit is not None:          # don't cache network failures
                self.cache[genus] = hit
                self.dirty = True
                self.save()              # keep progress if interrupted
        hit = self.cache.get(genus)
        if not hit:
            return "", f"family for '{genus}' not cached (run once with internet)"
        if not hit.get("family"):
            return "", f"GBIF found no family for genus '{genus}'"
        note = ""
        if hit.get("match") == "FUZZY":
            note = (f"genus '{genus}' fuzzy-matched to GBIF '{hit.get('gbif_genus')}' "
                    f"-> {hit['family']}; check spelling")
        return hit["family"], note

    def _query(self, genus: str) -> dict | None:
        q = urllib.parse.urlencode({"name": genus, "rank": "GENUS", "kingdom": "Animalia",
                                   "phylum": "Chordata"})
        req = urllib.request.Request(self.URL + q, headers={"User-Agent": "ccber-jar-labels/1.0"})
        d = None
        for attempt in range(4):
            try:
                with urllib.request.urlopen(req, timeout=15) as r:
                    d = json.load(r)
                break
            except Exception as e:  # offline, proxy, flaky connection
                err = e
                time.sleep(1.5 * (attempt + 1))
        if d is None:
            self.failures += 1
            if self.failures >= 5:
                self.online = False
                print(f"  GBIF unreachable ({err}); using cache only", file=sys.stderr)
            return None
        self.failures = 0
        ok = (d.get("kingdom") == "Animalia" and d.get("family")
              and d.get("matchType") in ("EXACT", "FUZZY") and d.get("confidence", 0) >= 80)
        return {"family": d.get("family", "") if ok else "",
                "match": d.get("matchType", ""), "gbif_genus": d.get("genus", "")}

    def save(self):
        if self.dirty:
            self.path.write_text(json.dumps(dict(sorted(self.cache.items())), indent=1))
            self.dirty = False


# ==========================================================================
# Workbook reading
# ==========================================================================
@dataclass
class Jar:
    sheet: str
    row: int
    jar_id: str
    sciname: str
    family: str
    header: str
    locality: str
    collector: str
    preservative: str
    catalogs: list[str]
    has_collector_col: bool
    problems: list[str] = field(default_factory=list)
    n_specimens: int = 0
    size_raw: str = ""


def find_header(ws, scan=10):
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=scan, values_only=True), 1):
        vals = [clean(v).lower() for v in row]
        idx = {}
        for key, names in COL.items():
            for n in names:
                if n in vals:
                    idx[key] = vals.index(n)
                    break
        if "catalog" in idx and "sciname" in idx:
            return r, idx
    return None, None


def read_jars(path: Path, sheets, jar_filter, max_expand, expand, families: FamilyLookup | None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    jars: list[Jar] = []
    skipped: list[dict] = []
    for ws in wb.worksheets:
        if sheets and ws.title not in sheets:
            continue
        hdr_row, idx = find_header(ws)
        if hdr_row is None:
            continue

        def get(row, key):
            i = idx.get(key)
            return clean(row[i]) if i is not None and i < len(row) else ""

        for rnum, row in enumerate(ws.iter_rows(min_row=hdr_row + 1, values_only=True),
                                   start=hdr_row + 1):
            if not any(clean(v) for v in row) or get(row, "shelf").upper() == "KEY":
                continue
            jar_id = get(row, "jar")
            if jar_filter and jar_id.upper() not in jar_filter:
                continue

            sci = normalize_sciname(get(row, "sciname"))
            cat_raw = get(row, "catalog")
            notes = []
            if not looks_like_catalog(cat_raw):
                for alt in ("side", "top"):
                    if looks_like_catalog(get(row, alt)):
                        cat_raw = get(row, alt)
                        notes.append(f"'Catalog #s' empty; used {alt.title()} Catalog")
                        break
            catalogs, problems = parse_catalog(cat_raw, max_expand, expand) if cat_raw else ([], [])

            if not sci and not catalogs:
                skipped.append({"sheet": ws.title, "row": rnum, "jar": jar_id,
                                "issue": "skipped: no scientific name and no catalog numbers"})
                continue

            family = get(row, "family")
            if not family and families and sci:
                family, note = families.get(sci)
                if note:
                    notes.append(note)

            jar = Jar(ws.title, rnum, jar_id, sci, family,
                      geography_header(get(row, "state"), get(row, "locality")),
                      locality_text(get(row, "locality")), get(row, "collector"),
                      get(row, "preservative"), catalogs, "collector" in idx,
                      notes + problems)
            if not sci:
                jar.problems.append("no scientific name")
            if not catalogs:
                jar.problems.append("no catalog numbers")
            if not jar.header:
                jar.problems.append("no usable state/country")
            if not jar_id:
                jar.problems.append("no Jar #")
            jar.size_raw = get(row, "size")
            jars.append(jar)
    return jars, skipped


# ==========================================================================
# Layout
# ==========================================================================
@dataclass
class Style:
    scale: float = 1.0
    footer: list[str] = field(default_factory=lambda: list(DEFAULT_FOOTER))
    margin_x: float = 0.5 * cm
    margin_y: float = 0.3 * cm
    valign: str = "center"
    show_preservative: bool = False

    # sizes from Label_format_example.docx (16 / 14 / 12 pt); rows ~10.5 pt
    def s(self, pt): return pt * self.scale
    @property
    def name_size(self): return self.s(16)
    @property
    def geo_size(self): return self.s(14)
    @property
    def head_size(self): return self.s(12)
    @property
    def row_size(self): return self.s(10.5)
    @property
    def row_pitch(self): return self.row_size * 1.4
    @property
    def footer_size(self): return self.s(12)
    thick: float = 1.5   # pt, matches the docx border weights
    thin: float = 0.75


def name_runs(name: str):
    """(text, font) runs: italic name, but 'sp.', 'cf.' etc. upright."""
    return [((" " if i else "") + w, FONT_REG if w.lower() in ROMAN_TOKENS else FONT_ITAL)
            for i, w in enumerate(name.split(" "))]


def runs_width(runs, size):
    return sum(stringWidth(t, f, size) for t, f in runs)


def fit_size(text, font, size, max_w, min_size):
    while size > min_size and stringWidth(text, font, size) > max_w:
        size -= 0.25
    return size


@dataclass
class Column:
    title: str
    width: float
    values: list[str] | str   # per-row list, or one value repeated
    size: float


@dataclass
class Layout:
    length: float
    height: float
    content_h: float
    columns: list[Column]
    group_w: float
    group_gap: float
    groups: int
    rows_per_group: int
    family_text: str
    jar_text: str
    natural_length: float = 0.0   # length the content needs
    fits_height: bool = True


def vertical_fixed(st: Style) -> float:
    """Height of everything except the specimen rows."""
    return (st.name_size * 1.3 + st.geo_size * 1.3 + 4          # title lines
            + st.head_size * 1.6 + 3                            # column headings
            + 4                                                 # space above bottom rule
            + len(st.footer) * st.footer_size * 1.2 + 5)        # footer


def plan(jar: Jar, st: Style, paper_h: float, fixed_len: float | None) -> Layout:
    avail_h = paper_h - 2 * st.margin_y
    fit_rows = max(int((avail_h - vertical_fixed(st)) // st.row_pitch), 1)
    n = max(len(jar.catalogs), 1)
    groups = math.ceil(n / fit_rows)
    rows = math.ceil(n / groups)

    cats = jar.catalogs or ["—"]
    cols = [Column("Catalog #", 0, cats, st.row_size)]
    if jar.locality:
        size = fit_size(jar.locality, FONT_REG, st.row_size, 8 * cm, st.s(7.5))
        cols.append(Column("Locality", 0, jar.locality, size))
    if jar.has_collector_col:
        size = fit_size(jar.collector, FONT_REG, st.row_size, 5 * cm, st.s(7.5))
        cols.append(Column("Collector", 0, jar.collector, size))

    pads = [1.0 * cm, 1.4 * cm, 0.2 * cm]
    mins = [2.6 * cm, 0, 0]
    for i, c in enumerate(cols):
        vals = c.values if isinstance(c.values, list) else [c.values]
        w = max([stringWidth(c.title, FONT_REG, st.head_size)] +
                [stringWidth(v, FONT_REG, c.size) for v in vals])
        last = i == len(cols) - 1
        c.width = max(w + (0.4 * cm if last else pads[min(i, 1)]), mins[min(i, 2)])
    group_w = sum(c.width for c in cols)
    gap = 1.2 * cm
    table_w = groups * group_w + (groups - 1) * gap

    family_text = f"    ({jar.family})" if jar.family else ""
    jar_text = f"Jar#: {jar.jar_id}" if jar.jar_id else ""
    title_w = (runs_width(name_runs(jar.sciname), st.name_size)
               + stringWidth(family_text, FONT_REG, st.name_size)
               + (stringWidth(jar_text, FONT_REG, st.name_size) + 1.2 * cm if jar_text else 0))
    geo_w = stringWidth(jar.header, FONT_REG, st.geo_size)
    foot_w = max((stringWidth(t, FONT_REG, st.footer_size) for t in st.footer), default=0) + 1.0 * cm
    if st.show_preservative and jar.preservative:
        foot_w += 2 * (stringWidth(jar.preservative, FONT_REG, st.s(8)) + 0.5 * cm)

    content_w = max(table_w, title_w, geo_w, foot_w)
    length = content_w + 2 * st.margin_x
    natural = length
    if fixed_len:
        length = max(length, fixed_len)

    content_h = vertical_fixed(st) + rows * st.row_pitch
    return Layout(length, paper_h, content_h, cols, group_w, gap, groups, rows,
                  family_text, jar_text, natural, content_h <= avail_h + 0.5)


# ==========================================================================
# Jar sizes
# ==========================================================================
# Thermo Scientific wide-mouth, tall-profile clear glass jars (catalog family
# 121-/221-/321-xxxx; the prefix is the cleaning level, the suffix the size).
# Diameter and height in mm are from thermofisher.com product pages, except
# the 500 mL jar, whose dimensions are not published there: ESTIMATED -- measure
# one and correct the numbers below.  Add your own jars here, or pass
# --jar-size DIAMETERxHEIGHT (mm), e.g. --jar-size 70x150.
@dataclass
class JarSpec:
    key: str
    description: str
    catalog: str
    diameter_mm: float
    height_mm: float
    estimated: bool = False


JAR_SIZES = {j.key: j for j in [
    JarSpec("125ml", "125 mL (4 oz) tall wide-mouth", "121-0125", 51, 102),
    JarSpec("250ml", "250 mL (8 oz) tall wide-mouth", "121-0250", 62, 127),
    JarSpec("500ml", "500 mL (16 oz) tall wide-mouth", "121-0500", 76, 150, estimated=True),
    JarSpec("950ml", "950 mL (32 oz) tall wide-mouth", "121-1000", 95, 170),
    JarSpec("2000ml", "2 L (64 oz) tall wide-mouth", "221-2000", 124, 214),
    JarSpec("4000ml", "4 L (128 oz / 1 gal) tall wide-mouth", "121-4000", 133, 253),
]}

# Other names accepted on the command line or in the sheet's "Jar Size" column.
# The small/medium/large words are a GUESS at what the data-entry team meant --
# edit them to match your shelves.
JAR_ALIASES = {
    "4oz": "125ml", "8oz": "250ml", "16oz": "500ml", "pint": "500ml",
    "32oz": "950ml", "1l": "950ml", "quart": "950ml",
    "64oz": "2000ml", "2l": "2000ml", "halfgallon": "2000ml",
    "128oz": "4000ml", "4l": "4000ml", "gallon": "4000ml", "1gallon": "4000ml",
    "xxsmall": "125ml", "xsmall": "125ml", "small": "250ml", "medium": "500ml",
    "large": "950ml", "xlarge": "2000ml", "xxlarge": "4000ml",
}
CATALOG_SUFFIX = {"0125": "125ml", "0250": "250ml", "0500": "500ml",
                  "1000": "950ml", "2000": "2000ml", "4000": "4000ml"}

# How a jar turns into a label
LABEL_HEIGHT_FRACTION = 0.6   # label height = 60% of jar height (clears shoulder & base)
FULL_WRAP_GAP_MM = 10         # never longer than circumference minus this


def resolve_jar_size(value: str) -> tuple[JarSpec | None, str]:
    """'small', '8 oz', '250 mL', '121-0250', '70x150' -> JarSpec.  Returns (spec, note)."""
    raw = clean(value)
    if not raw:
        return None, ""
    v = re.sub(r"[\s.\-_]", "", raw.lower()).replace("ounces", "oz").replace("ounce", "oz")
    if (m := re.fullmatch(r"(\d+(?:\.\d+)?)[x\u00d7*](\d+(?:\.\d+)?)(mm)?", v)):
        d, h = float(m.group(1)), float(m.group(2))
        return JarSpec(f"{d:g}x{h:g}", f"custom {d:g} x {h:g} mm jar", "", d, h), ""
    if (m := re.search(r"[123]21-?(\d{4})", raw)) and m.group(1) in CATALOG_SUFFIX:
        return JAR_SIZES[CATALOG_SUFFIX[m.group(1)]], ""
    if v in JAR_SIZES:
        return JAR_SIZES[v], ""
    if v in JAR_ALIASES:
        return JAR_SIZES[JAR_ALIASES[v]], ""
    hit = difflib.get_close_matches(v, list(JAR_SIZES) + list(JAR_ALIASES), n=1, cutoff=0.8)
    if hit:
        key = JAR_ALIASES.get(hit[0], hit[0])
        return JAR_SIZES[key], f"jar size '{raw}' read as '{hit[0]}' ({key})"
    return None, f"jar size '{raw}' not recognised; label sized to content"


def jar_label_box(spec: JarSpec, paper_h: float, wrap: float) -> tuple[float, float, float]:
    """(label height, target length, max length) in points for a jar."""
    circ = math.pi * spec.diameter_mm
    height = min(paper_h, spec.height_mm * LABEL_HEIGHT_FRACTION / 10 * cm)
    return height, wrap * circ / 10 * cm, (circ - FULL_WRAP_GAP_MM) / 10 * cm


def fit_to_box(jar: Jar, st: Style, box_h: float, box_len: float, min_scale: float):
    """Largest type scale (<= st.scale) whose layout fits box_h x box_len."""
    scale = st.scale
    while True:
        s2 = replace(st, scale=scale)
        L = plan(jar, s2, box_h, box_len)
        if (L.fits_height and L.natural_length <= box_len + 0.5) or scale <= min_scale + 1e-9:
            return L, s2
        scale = max(min_scale, round(scale - 0.025, 3))


def print_jar_table(paper_h: float, wrap: float):
    print(f"{'size':8} {'catalog':9} {'diam':>5} {'height':>6}   label (h x length)  description")
    for j in JAR_SIZES.values():
        h, length, _ = jar_label_box(j, paper_h, wrap)
        est = "  * dimensions estimated" if j.estimated else ""
        print(f"{j.key:8} {j.catalog:9} {j.diameter_mm:4g}mm {j.height_mm:4g}mm   "
              f"{h / cm:4.1f} x {length / cm:4.1f} cm     {j.description}{est}")
    print("\naliases:", ", ".join(f"{k}={v}" for k, v in JAR_ALIASES.items()))


# ==========================================================================
# Drawing
# ==========================================================================
def draw_label(c: canvas.Canvas, jar: Jar, st: Style, L: Layout, rotate: bool,
               page_h: float | None = None, cut_line: bool = True):
    """page_h = roll width.  When the label is narrower than the roll it sits
    against the top edge and a dashed cut line marks its bottom edge."""
    page_h = max(page_h or L.height, L.height)
    if rotate:  # page width = roll width, page height = label length
        c.setPageSize((page_h, L.length))
        c.translate(page_h, 0)
        c.rotate(90)
    else:
        c.setPageSize((L.length, page_h))
    if page_h - L.height > 0.1 * cm:
        if cut_line:
            c.saveState()
            c.setLineWidth(0.5)
            c.setDash(3, 3)
            c.line(0, page_h - L.height, L.length, page_h - L.height)
            c.restoreState()
        c.translate(0, page_h - L.height)

    x0, x1 = st.margin_x, L.length - st.margin_x
    if st.valign == "top":
        y = L.height - st.margin_y
    else:
        y = L.height - max(st.margin_y, (L.height - L.content_h) / 2)

    # --- line 1: name (family) .......... Jar#: H1175
    y -= st.name_size
    x = x0
    for text, font in name_runs(jar.sciname):
        c.setFont(font, st.name_size)
        c.drawString(x, y, text)
        x += stringWidth(text, font, st.name_size)
    if L.family_text:
        c.setFont(FONT_REG, st.name_size)
        c.drawString(x, y, L.family_text)
        x += stringWidth(L.family_text, FONT_REG, st.name_size)
    if L.jar_text:
        c.setFont(FONT_REG, st.name_size)
        jw = stringWidth(L.jar_text, FONT_REG, st.name_size)
        # as in the example: aligned with the last column when there is room
        jx = x0 + sum(col.width for col in L.columns[:-1]) if len(L.columns) > 1 else x1 - jw
        if L.groups > 1 or jx < x + 1.2 * cm or jx + jw > x1:
            jx = x1 - jw
        c.drawString(jx, y, L.jar_text)
    y -= st.name_size * 0.3

    # --- line 2: geography
    y -= st.geo_size
    c.setFont(FONT_REG, st.geo_size)
    c.drawString(x0, y, jar.header)
    y -= st.geo_size * 0.3 + 4

    c.setLineWidth(st.thick)
    c.line(x0, y, x1, y)

    # --- column headings
    y -= st.head_size * 1.15
    c.setFont(FONT_REG, st.head_size)
    for g in range(L.groups):
        cx = x0 + g * (L.group_w + L.group_gap)
        for col in L.columns:
            c.drawString(cx, y, col.title)
            cx += col.width
    y -= st.head_size * 0.45 + 3
    c.setLineWidth(st.thin)
    c.line(x0, y, x1, y)

    # --- rows
    n = len(L.columns[0].values)
    for i in range(n):
        g, r = divmod(i, L.rows_per_group)
        cx = x0 + g * (L.group_w + L.group_gap)
        yy = y - (r + 1) * st.row_pitch + (st.row_pitch - st.row_size) * 0.45
        for col in L.columns:
            val = col.values[i] if isinstance(col.values, list) else col.values
            c.setFont(FONT_REG, col.size)
            c.drawString(cx, yy, val)
            cx += col.width
    y -= L.rows_per_group * st.row_pitch + 4
    c.setLineWidth(st.thick)
    c.line(x0, y, x1, y)

    # --- footer, centred
    y -= 3
    c.setFont(FONT_REG, st.footer_size)
    for line in st.footer:
        y -= st.footer_size * 1.2
        c.drawCentredString((x0 + x1) / 2, y + st.footer_size * 0.2, line)
    if st.show_preservative and jar.preservative:
        c.setFont(FONT_REG, st.s(8))
        c.drawRightString(x1, y + st.footer_size * 0.2, jar.preservative)
    c.showPage()


# ==========================================================================
# CLI
# ==========================================================================
def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("xlsx", type=Path, nargs="?", help="jar inventory workbook (.xlsx)")
    p.add_argument("-o", "--output", type=Path, default=Path("jar_labels.pdf"))
    p.add_argument("--sheet", nargs="+", help="only these sheets (default: every sheet with "
                   "'Catalog #s' and 'Scientific Name' headers)")
    p.add_argument("--jar", nargs="+", help="only these Jar #s, e.g. H1874 H1875")
    p.add_argument("--limit", type=int, help="stop after N labels (test prints)")
    p.add_argument("--paper-width", type=float, default=10.0,
                   help="roll width in cm = label height (default 10)")
    p.add_argument("--length", type=float, help="fixed label length in cm (default: fit content)")
    p.add_argument("--jar-size", metavar="SIZE",
                   help="size every label for this jar: a size key (125ml 250ml 500ml 950ml "
                        "2000ml 4000ml), an alias (8oz, small, quart...), a Thermo catalog # "
                        "(121-0250) or DIAMETERxHEIGHT in mm (70x150).  Without it, each row's "
                        "'Jar Size' cell is used, and rows with no size are sized to content.")
    p.add_argument("--list-jar-sizes", action="store_true", help="show the jar table and exit")
    p.add_argument("--wrap", type=float, default=0.75,
                   help="fraction of the jar's circumference the label wraps (default 0.75)")
    p.add_argument("--min-scale", type=float, default=0.6,
                   help="smallest type scale used to squeeze a label onto its jar (default 0.6)")
    p.add_argument("--no-cut-line", action="store_true",
                   help="omit the dashed trim line under labels narrower than the roll")
    p.add_argument("--max-length", type=float, default=30.0,
                   help="flag labels longer than this in the review file (cm; default 30)")
    p.add_argument("--scale", type=float, default=1.0, help="scale all type, e.g. 0.9")
    p.add_argument("--valign", choices=["center", "top"], default="center",
                   help="where the label sits across the roll")
    p.add_argument("--footer", nargs="+", metavar="LINE", help="replace the footer lines")
    p.add_argument("--font-regular", help="TrueType file for regular text (e.g. Avenir Book .ttf)")
    p.add_argument("--font-italic", help="TrueType file for the italic scientific name")
    p.add_argument("--no-gbif", action="store_true",
                   help="don't query GBIF for families; use family_cache.json only")
    p.add_argument("--family-cache", type=Path,
                   default=Path(__file__).with_name("family_cache.json"))
    p.add_argument("--no-family", action="store_true", help="leave the family off")
    p.add_argument("--min-run", type=int, default=2,
                   help="consecutive numbers needed before they print as a range "
                        "(default 2: 40748, 40749 -> '40748 - 40749')")
    p.add_argument("--no-ranges", action="store_true",
                   help="print every catalog number on its own row instead of as ranges")
    p.add_argument("--max-expand", type=int, default=300,
                   help="spreadsheet ranges spanning more numbers than this are treated as "
                        "suspect: printed as written and flagged")
    p.add_argument("--preservative", action="store_true", help="print preservative bottom-right")
    p.add_argument("--rotate", action="store_true",
                   help="rotate pages 90° (for drivers set to portrait)")
    a = p.parse_args(argv)
    if a.list_jar_sizes:
        print_jar_table(a.paper_width * cm, a.wrap)
        return
    if not a.xlsx:
        p.error("the workbook path is required")
    forced = None
    if a.jar_size:
        forced, note = resolve_jar_size(a.jar_size)
        if forced is None:
            p.error(f"unknown --jar-size '{a.jar_size}' (see --list-jar-sizes)")
        if note:
            print(note)

    print("Font:", setup_fonts(a.font_regular, a.font_italic))
    families = None if a.no_family else FamilyLookup(a.family_cache, online=not a.no_gbif)
    jar_filter = {j.upper() for j in a.jar} if a.jar else None
    jars, skipped = read_jars(a.xlsx, a.sheet, jar_filter, a.max_expand, True, families)
    if families:
        families.save()
    if a.limit:
        jars = jars[: a.limit]
    if not jars:
        sys.exit("No jars matched.")

    st = Style(scale=a.scale, valign=a.valign, show_preservative=a.preservative)
    if a.footer:
        st.footer = a.footer
    paper_h = a.paper_width * cm
    fixed_len = a.length * cm if a.length else None

    c = canvas.Canvas(str(a.output))
    c.setTitle("Jar labels")
    max_len = a.max_length * cm if a.max_length else None
    sizes_used: dict[str, int] = defaultdict(int)
    for jar in jars:
        jar.n_specimens = len(jar.catalogs)
        if not a.no_ranges:
            jar.catalogs = collapse_runs(jar.catalogs, a.min_run)
        spec, note = (forced, "") if forced else resolve_jar_size(jar.size_raw)
        if note:
            jar.problems.append(note)
        if spec:
            if spec.estimated:
                jar.problems.append(f"{spec.key} jar dimensions are estimated; check the label fits")
            box_h, box_len, wrap_max = jar_label_box(spec, paper_h, a.wrap)
            if fixed_len:
                box_len = fixed_len
            L, used = fit_to_box(jar, st, box_h, box_len, a.min_scale)
            if not L.fits_height or L.natural_length > box_len + 0.5:
                L = plan(jar, used, box_h, max(box_len, L.natural_length))
                over = L.length > wrap_max
                jar.problems.append(
                    f"doesn't fit a {spec.key} jar label ({box_h / cm:.1f} x {box_len / cm:.1f} cm) "
                    f"even at {used.scale:.0%} type; printed "
                    + (f"{L.length / cm:.1f} cm long" if L.fits_height else "with rows overflowing")
                    + ("; LONGER THAN THE JAR'S CIRCUMFERENCE - use a bigger jar" if over else ""))
            draw_label(c, jar, used, L, a.rotate, paper_h, not a.no_cut_line)
        else:
            L = plan(jar, st, paper_h, fixed_len)
            if fixed_len and L.natural_length > fixed_len + 0.5:
                jar.problems.append(f"content needs {L.natural_length / cm:.1f} cm; --length is "
                                    f"{a.length:g} cm (label made longer)")
            if not L.fits_height:
                jar.problems.append("label taller than paper; try --scale 0.9")
            if max_len and L.length > max_len:
                jar.problems.append(f"label is {L.length / cm:.0f} cm long (> --max-length "
                                    f"{a.max_length:g} cm)")
            draw_label(c, jar, st, L, a.rotate)
        sizes_used[spec.key if spec else "fit to content"] += 1
    c.save()

    review = a.output.with_name(a.output.stem + "_review.csv")
    with review.open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["page", "sheet", "row", "jar", "scientific_name", "n_specimens", "issue"])
        for page, jar in enumerate(jars, 1):
            for issue in jar.problems:
                w.writerow([page, jar.sheet, jar.row, jar.jar_id, jar.sciname,
                            jar.n_specimens, issue])
        for s in skipped:
            w.writerow(["", s["sheet"], s["row"], s["jar"], "", 0, s["issue"]])

    print(f"{len(jars)} labels -> {a.output}")
    print("  by jar size: " + ", ".join(f"{k} {v}" for k, v in sizes_used.items()))
    print(f"{sum(1 for j in jars if j.problems)} labels flagged, "
          f"{len(skipped)} rows skipped -> {review}")


if __name__ == "__main__":
    main()
